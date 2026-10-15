"""Servicio `PlantillaReporte` — actas oficiales XLSX sobre plantillas
corporativas (cambio 2026-10-01-reportes-consumo-management).

REGLAS (REQ-P2-REP-001/002/006/007):
  * El backend NUNCA genera un workbook desde cero: SIEMPRE carga una de las
    dos plantillas corporativas de `backend/app/templates/` con openpyxl e
    INYECTA datos en sus celdas, conservando estilos, celdas combinadas y la
    imagen `company_logobrand.png` ya embebida (requiere Pillow).
  * Las rutas de plantilla son CONSTANTES del servidor: ningún dato de
    usuario participa en la construcción de rutas (anti path traversal).
  * Si los ítems exceden las filas de datos de la plantilla, se insertan
    filas copiando el estilo de una fila de datos y se DESPLAZAN
    manualmente todo lo que queda debajo (celdas, merges, alturas de fila)
    y las imágenes ancladas en/bajo el punto de inserción: `insert_rows` de
    openpyxl no mueve merges, alturas ni anclas de imagen.
  * Todo texto inyectado se sanitiza contra inyección de fórmulas (OWASP):
    openpyxl escribe como FÓRMULA cualquier cadena que empiece por `=`.
  * CERO aritmética de stock: las cantidades llegan ya calculadas por SQL
    (snapshot jsonb de las actas); aquí solo se transportan a celdas.
"""

import io
import re
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

FormatoActa = Literal["normal", "extenso"]
TipoActa = Literal["ENTREGA", "CONSUMO"]

MEDIA_TYPE_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# ── Constantes del servidor (jamás derivadas de input de usuario) ─────────
TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
_PLANTILLAS: dict[str, str] = {
    "normal": "1 ACTA DE REPORTE NORMAL.xlsx",
    "extenso": "2 ACTA DE REPORTE EXTENSO.xlsx",
}

TITULO_ACTA_CONSUMO = "ACTA DE CONSUMO"

_MESES = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
    "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)  # fmt: skip

_COL_CANT = 1  # A
_COL_DESC = 2  # B (combinada B:D)
_COL_DESC_FIN = 4  # D
_COL_UM = 5  # E


@dataclass(frozen=True)
class _Geometria:
    """Geometría de una plantilla (verificada contra los XLSX corporativos)."""

    primera_fila: int  # primera fila de datos
    ultima_fila: int  # última fila de datos (capacidad base)
    fila_estilo: int  # fila de datos cuyo estilo se copia al desbordar
    celda_titulo: str = "A2"
    celda_fecha: str = "A5"


_GEOMETRIA: dict[str, _Geometria] = {
    # Normal: datos 10-25, fila 26 separadora, bloque de firmas en 27.
    "normal": _Geometria(primera_fila=10, ultima_fila=25, fila_estilo=10),
    # Extenso: datos 9-28, sección COLABORADORES desde la fila 30.
    "extenso": _Geometria(primera_fila=9, ultima_fila=28, fila_estilo=10),
}


@dataclass(frozen=True)
class LineaActa:
    """Línea ya resuelta del acta: (cantidad, descripción, U.M.).
    `cantidad` es Decimal (RF-6: cantidades con hasta 3 decimales)."""

    cantidad: Decimal
    descripcion: str
    u_m: str | None


# ── Sanitización ──────────────────────────────────────────────────────────


def sanitizar_celda(valor):
    """Neutraliza la inyección de fórmulas (REQ-P2-REP-007).

    Todo texto cuyo contenido (sin espacios líderes) inicie con `=`, `+`,
    `-` o `@` se prefija con una comilla simple, de modo que openpyxl lo
    escriba como cadena y no como fórmula. Los números (int/float/Decimal)
    se devuelven tal cual (no son texto; openpyxl los escribe como celdas
    numéricas)."""
    if valor is None or isinstance(valor, (int, float, Decimal)):
        return valor
    texto = str(valor)
    if texto.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + texto
    return texto


def sanitizar_nombre_archivo(nombre: str) -> str:
    """Nombre de archivo seguro: solo `[A-Za-z0-9._-]`, sin separadores de
    ruta ni caracteres de control (anti path traversal / header injection)."""
    limpio = re.sub(r"[^A-Za-z0-9._-]", "_", nombre).strip("._")
    return limpio or "acta"


def nombre_archivo_acta(
    tipo: TipoActa, formato: FormatoActa, fecha: date
) -> str:
    """Nombre del adjunto, construido solo con constantes y la fecha."""
    return sanitizar_nombre_archivo(
        f"acta-{tipo.lower()}-{formato}-{fecha.isoformat()}.xlsx"
    )


def formatear_fecha(fecha: date) -> str:
    """`FECHA: 28 / SEPTIEMBRE / 2026` (formato de las plantillas)."""
    return f"FECHA: {fecha.day:02d} / {_MESES[fecha.month - 1]} / {fecha.year}"


# ── Desplazamiento de filas (insert_rows no mueve merges/alturas/imágenes) ─


def _desplazar_filas(ws: Worksheet, desde_fila: int, n: int) -> None:
    """Abre `n` filas vacías en `desde_fila`, desplazando hacia abajo todo lo
    que está en `desde_fila` o debajo: celdas (valor + estilo), celdas
    combinadas, alturas de fila e imágenes ancladas."""
    if n <= 0:
        return
    max_fila = ws.max_row
    max_col = ws.max_column

    # 1) Merges bajo el punto de inserción: se desanclan, se recuerdan y se
    #    vuelven a crear desplazados (move_range no los traslada).
    a_mover = [
        (r.min_col, r.min_row, r.max_col, r.max_row)
        for r in list(ws.merged_cells.ranges)
        if r.min_row >= desde_fila
    ]
    for min_col, min_row, max_c, max_r in a_mover:
        ws.unmerge_cells(
            start_row=min_row, start_column=min_col, end_row=max_r, end_column=max_c
        )

    # 2) Celdas (valor + estilo) del bloque inferior.
    if desde_fila <= max_fila:
        ws.move_range(
            f"A{desde_fila}:{ws.cell(row=max_fila, column=max_col).coordinate}",
            rows=n,
            cols=0,
            translate=False,
        )

    # 3) Alturas de fila (de abajo hacia arriba para no pisar valores).
    for fila in range(max_fila, desde_fila - 1, -1):
        altura = ws.row_dimensions[fila].height
        ws.row_dimensions[fila + n].height = altura
    for fila in range(desde_fila, desde_fila + n):
        ws.row_dimensions[fila].height = None

    # 4) Merges desplazados.
    for min_col, min_row, max_c, max_r in a_mover:
        ws.merge_cells(
            start_row=min_row + n,
            start_column=min_col,
            end_row=max_r + n,
            end_column=max_c,
        )

    # 5) Imágenes ancladas en o bajo el punto de inserción (el logo de las
    #    plantillas está en la fila 0, por lo que normalmente no se mueve;
    #    el re-anclado se mantiene para cualquier imagen posterior).
    for imagen in cast(Any, ws)._images:
        anchor = imagen.anchor
        marcador = getattr(anchor, "_from", None)
        if marcador is not None and marcador.row + 1 >= desde_fila:
            marcador.row += n
            destino = getattr(anchor, "to", None)
            if destino is not None:
                destino.row += n


def _insertar_filas_de_datos(
    ws: Worksheet, geo: _Geometria, desde_fila: int, n: int
) -> None:
    """Inserta `n` filas de datos en `desde_fila` copiando estilo, altura y
    combinación B:D de la fila de estilo de la plantilla."""
    altura_estilo = ws.row_dimensions[geo.fila_estilo].height
    _desplazar_filas(ws, desde_fila, n)
    for fila in range(desde_fila, desde_fila + n):
        for col in range(_COL_CANT, _COL_UM + 1):
            origen = ws.cell(row=geo.fila_estilo, column=col)
            destino = ws.cell(row=fila, column=col)
            if origen.has_style:
                cast(Any, destino)._style = copy(cast(Any, origen)._style)
        ws.merge_cells(
            start_row=fila,
            start_column=_COL_DESC,
            end_row=fila,
            end_column=_COL_DESC_FIN,
        )
        ws.row_dimensions[fila].height = altura_estilo


def _uniformar_estilo_fila(ws: Worksheet, geo: _Geometria, fila: int) -> None:
    """Las filas de datos de la plantilla Extenso no son uniformes (p. ej.
    filas 25, 27 y 28 con fuente 10 y sin centrado). Al escribir una fila se
    le aplica el estilo de la fila de datos de referencia para que todo el
    cuerpo del acta luzca igual; no se altera ninguna celda no usada."""
    if fila == geo.fila_estilo:
        return
    for col in range(_COL_CANT, _COL_UM + 1):
        origen = ws.cell(row=geo.fila_estilo, column=col)
        destino = ws.cell(row=fila, column=col)
        if origen.has_style:
            cast(Any, destino)._style = copy(cast(Any, origen)._style)


# ── API pública ───────────────────────────────────────────────────────────


def generar_acta(
    *,
    tipo: TipoActa,
    formato: FormatoActa,
    fecha: date,
    lineas: list[LineaActa],
) -> bytes:
    """Genera el XLSX del acta inyectando datos sobre la plantilla
    corporativa (nunca desde cero) y devuelve los bytes del archivo.

    * `formato`: 'normal' o 'extenso' (plantilla 1 o 2).
    * `tipo='CONSUMO'`: el título pasa a "ACTA DE CONSUMO".
    * Desborde: si `len(lineas)` excede la capacidad de la plantilla se
      insertan filas (estilo copiado, merges/alturas/imágenes desplazados).
    """
    if formato not in _PLANTILLAS:  # defensa: formato validado por el router
        raise ValueError(f"Formato de acta no soportado: {formato!r}")
    geo = _GEOMETRIA[formato]
    wb = load_workbook(TEMPLATES_DIR / _PLANTILLAS[formato])
    ws = cast(Worksheet, wb.active)

    capacidad = geo.ultima_fila - geo.primera_fila + 1
    exceso = len(lineas) - capacidad
    if exceso > 0:
        _insertar_filas_de_datos(ws, geo, geo.ultima_fila + 1, exceso)

    if tipo == "CONSUMO":
        ws[geo.celda_titulo] = TITULO_ACTA_CONSUMO
    ws[geo.celda_fecha] = formatear_fecha(fecha)

    for i, linea in enumerate(lineas):
        fila = geo.primera_fila + i
        _uniformar_estilo_fila(ws, geo, fila)
        ws.cell(row=fila, column=_COL_CANT, value=sanitizar_celda(linea.cantidad))
        ws.cell(row=fila, column=_COL_DESC, value=sanitizar_celda(linea.descripcion))
        ws.cell(row=fila, column=_COL_UM, value=sanitizar_celda(linea.u_m))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
