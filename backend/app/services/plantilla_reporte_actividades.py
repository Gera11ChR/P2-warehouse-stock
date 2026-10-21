"""Servicio `PlantillaReporteActividades` — reporte oficial XLSX de
ACTIVIDADES por Unidad sobre la plantilla corporativa
«REPORTE DE ACTIVIDADES.xlsx» (paquete 2026-10-actividades-reportes,
Etapa 2/3; REQ-REP-006/007/008/009/010).

REGLAS:
  * El backend NUNCA genera un workbook desde cero: carga la plantilla
    corporativa de `backend/app/templates/` (hoja `401-450`) con openpyxl e
    INYECTA datos en sus celdas, conservando estilos, celdas combinadas,
    alturas y el logo corporativo (REQ-REP-006). La plantilla en disco NUNCA
    se escribe: solo `load_workbook` en memoria.
  * Una hoja por folio, titulada con el folio, en orden de folio ascendente;
    la hoja residual `Sheet1` de la plantilla no se incluye y TODAS las
    hojas incluyen el logo (REQ-REP-009).
  * Logo: se captura UNA sola vez (`base._images[0].ref.getvalue()`, un
    BytesIO) ANTES de copiar o guardar; cada copia recibe su propia imagen
    con un `OneCellAnchor` nuevo. PROHIBIDO `Image._data()` sobre la imagen
    de la plantilla: es de un solo uso, cierra el flujo y rompe `wb.save`.
  * Desborde (REQ-REP-008): más de 10 materiales, más de 4 actividades
    realizadas o comentarios más largos que su área insertan filas copiando
    estilo, altura y celdas combinadas de la fila de referencia y desplazan
    lo inferior (celdas, merges, alturas e imágenes) con `_desplazar_filas`.
  * Alturas por heurística determinista: líneas estimadas × 13,2 pt por
    fila, repartidas en el área de comentarios, con tope de 409,5 pt por
    fila (límite de Excel).
  * FECHA como TEXTO `DD / MES / AAAA` en la zona de negocio, sin prefijo
    «FECHA:» (la etiqueta ya está en A8) (REQ-REP-007).
  * Todo texto inyectado pasa por `sanitizar_celda` (neutraliza fórmulas,
    REQ-REP-010) y por `ILLEGAL_CHARACTERS_RE`; los números se escriben como
    números.
  * CERO aritmética de negocio: cantidades y valor unitario del snapshot se
    transportan tal cual (REQ-REP-007).
"""

import io
from collections.abc import Sequence
from copy import copy
from io import BytesIO
from math import ceil
from typing import Any, cast
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.actividad import ActividadOut
from app.services.plantilla_reporte import (
    TEMPLATES_DIR,
    _MESES,
    _desplazar_filas,
    sanitizar_celda,
    sanitizar_nombre_archivo,
)

# ── Constantes del servidor (jamás derivadas de input de usuario) ─────────
PLANTILLA_ACTIVIDADES = "REPORTE DE ACTIVIDADES.xlsx"
HOJA_PLANTILLA = "401-450"
HOJA_RESIDUAL = "Sheet1"

# A13/I13 son 'General' en la plantilla: el formato de VALOR UNITARIO se
# fija explícitamente en las celdas I de datos (plan §3.2/§3.3 paso 5).
FORMATO_VALOR_UNITARIO = "#,##0.00#"

# ── Geometría de la plantilla (verificada con openpyxl, plan §3.2) ────────
CELDA_FECHA = "B8"  # combinada B8:C8; recibe TEXTO
CELDA_VEHICULO = "E8"  # combinada E8:F8
CELDA_NOMBRE = "B9"  # combinada B9:F9 (clientes)
CELDA_FOLIO = "I9"  # combinada I9:J9 (texto 'FOLIO: R0000')
CELDA_DIRECCION = "B10"  # combinada B10:F10

FILA_MATERIAL_INICIO = 13  # filas 13-22, cada una con C:H combinada
MAX_MATERIALES = 10
FILA_ACTIVIDAD_INICIO = 25  # filas 25-28, cada una con A:I combinada
MAX_ACTIVIDADES = 4
FILA_COMENTARIO_INICIO = 32  # etiqueta A32:B34, valor C32:I34
FILAS_COMENTARIO = 3

FILA_INSERCION_MATERIALES = 23
FILA_INSERCION_ACTIVIDADES = 29
FILA_INSERCION_COMENTARIOS = 35

# ── Ancla y tamaño del logo (paridad con la imagen de la plantilla) ───────
ANCLA_COL = 6
ANCLA_FILA = 0
ANCLA_COL_OFF = 560070
ANCLA_ROW_OFF = 22860
ANCLA_EXT_CX = 2266950
ANCLA_EXT_CY = 876300

# ── Heurística de alturas (plan §3.3 paso 6) ──────────────────────────────
ALTURA_BASE = 13.2
TOPE_ALTURA = 409.5
ANCHO_DEFECTO = 8.43


def _limpiar_texto(valor: str) -> str:
    """Texto listo para XLSX: elimina caracteres ilegales del formato y
    después neutraliza fórmulas (REQ-REP-010).

    El orden importa: si se neutralizara primero, un prefijo de control
    (p. ej. ``\\x01=1+1``) ocultaría el ``=`` a ``sanitizar_celda`` y, al
    quitar el carácter ilegal, quedaría una fórmula real."""
    return str(sanitizar_celda(ILLEGAL_CHARACTERS_RE.sub("", str(valor))))


def _ancho_combinado(ws: Worksheet, col_inicio: int, col_fin: int) -> float:
    """Ancho real (unidades de columna) de un rango de columnas."""
    total = 0.0
    for col in range(col_inicio, col_fin + 1):
        dimension = ws.column_dimensions.get(get_column_letter(col))
        if dimension is not None:
            total += dimension.width
        else:
            total += ws.sheet_format.defaultColWidth or ANCHO_DEFECTO
    return total


def _caracteres_por_linea(ws: Worksheet, col_inicio: int, col_fin: int) -> int:
    return max(1, int(_ancho_combinado(ws, col_inicio, col_fin) * 0.95))


def _lineas_texto(texto: str, cpl: int) -> int:
    """Líneas estimadas de un texto a `cpl` caracteres por línea: una por
    párrafo (separado por `\\n`), mínimo 1."""
    return sum(max(1, ceil(len(parrafo) / cpl)) for parrafo in texto.split("\n"))


def _alto_para_lineas(lineas: int) -> float:
    return min(TOPE_ALTURA, max(ALTURA_BASE, lineas * ALTURA_BASE))


def _copiar_estilo_fila(
    ws: Worksheet, fila_origen: int, fila_destino: int, col_inicio: int, col_fin: int
) -> None:
    """Copia el estilo de `col_inicio..col_fin` de la fila de referencia
    (patrón de `_insertar_filas_de_datos` de `plantilla_reporte`)."""
    for col in range(col_inicio, col_fin + 1):
        origen = ws.cell(row=fila_origen, column=col)
        destino = ws.cell(row=fila_destino, column=col)
        if origen.has_style:
            cast(Any, destino)._style = copy(cast(Any, origen)._style)


def _insertar_filas_materiales(ws: Worksheet, fila_insercion: int, n: int) -> None:
    """Abre `n` filas de materiales copiando estilo, altura y combinación
    C:H de la fila 13 (nunca desplazada)."""
    altura = ws.row_dimensions[FILA_MATERIAL_INICIO].height
    _desplazar_filas(ws, fila_insercion, n)
    for fila in range(fila_insercion, fila_insercion + n):
        _copiar_estilo_fila(ws, FILA_MATERIAL_INICIO, fila, 1, 9)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=8)
        ws.row_dimensions[fila].height = altura


def _insertar_filas_actividades(
    ws: Worksheet, fila_insercion: int, n: int, fila_estilo: int
) -> None:
    """Abre `n` filas de actividades copiando estilo, altura y combinación
    A:I de la primera fila de actividades (ya desplazada, `25 + n_m`)."""
    altura = ws.row_dimensions[fila_estilo].height
    _desplazar_filas(ws, fila_insercion, n)
    for fila in range(fila_insercion, fila_insercion + n):
        _copiar_estilo_fila(ws, fila_estilo, fila, 1, 9)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
        ws.row_dimensions[fila].height = altura


def _insertar_filas_comentarios(
    ws: Worksheet, fila_insercion: int, n: int, fila_estilo: int
) -> None:
    """Abre `n` filas en el área de comentarios copiando el estilo de la
    fila del valor (ya desplazada, `32 + n_m + n_a`)."""
    _desplazar_filas(ws, fila_insercion, n)
    for fila in range(fila_insercion, fila_insercion + n):
        _copiar_estilo_fila(ws, fila_estilo, fila, 1, 9)
        ws.row_dimensions[fila].height = ALTURA_BASE


def _extender_merges_comentarios(
    ws: Worksheet, fila_inicio: int, fila_fin: int
) -> None:
    """Rehace las combinaciones del área de comentarios abarcando las filas
    insertadas: etiqueta A:B y valor C:I desde `fila_inicio` a `fila_fin`."""
    ws.unmerge_cells(
        start_row=fila_inicio, start_column=1, end_row=fila_inicio + 2, end_column=2
    )
    ws.merge_cells(start_row=fila_inicio, start_column=1, end_row=fila_fin, end_column=2)
    ws.unmerge_cells(
        start_row=fila_inicio, start_column=3, end_row=fila_inicio + 2, end_column=9
    )
    ws.merge_cells(start_row=fila_inicio, start_column=3, end_row=fila_fin, end_column=9)


def _incrustar_logo(hoja: Worksheet, logo_bytes: bytes, ancho: int, alto: int) -> None:
    """Re-incrusta el logo en una hoja copiada con ancla propia:
    `copy_worksheet` no copia imágenes y la de la plantilla es de un solo
    uso."""
    imagen = Image(BytesIO(logo_bytes))
    imagen.width = ancho
    imagen.height = alto
    imagen.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=ANCLA_COL, row=ANCLA_FILA, colOff=ANCLA_COL_OFF, rowOff=ANCLA_ROW_OFF
        ),
        ext=XDRPositiveSize2D(cx=ANCLA_EXT_CX, cy=ANCLA_EXT_CY),
    )
    hoja.add_image(imagen)


def _llenar_hoja(ws: Worksheet, actividad: ActividadOut, zona: ZoneInfo) -> None:
    """Inyecta los datos de una actividad en su hoja (plan §3.3, pasos 4-6).

    Los conteos de desborde se calculan ANTES de insertar; las inserciones
    van de arriba hacia abajo con offsets acumulados, de modo que las
    coordenadas de escritura finales son: materiales 13..12+len, actividades
    25+n_m..24+n_m+len y comentarios en (col 3, fila 32+n_m+n_a).
    """
    materiales = actividad.materiales
    realizadas = actividad.actividades
    comentarios = actividad.comentarios

    n_m = max(0, len(materiales) - MAX_MATERIALES)
    n_a = max(0, len(realizadas) - MAX_ACTIVIDADES)

    cpl_material = _caracteres_por_linea(ws, 3, 8)
    cpl_actividad = _caracteres_por_linea(ws, 1, 9)
    cpl_comentario = _caracteres_por_linea(ws, 3, 9)

    n_c = 0
    per_row_comentario = ALTURA_BASE
    if comentarios:
        total_pt = _lineas_texto(comentarios, cpl_comentario) * ALTURA_BASE
        filas_totales = FILAS_COMENTARIO
        while filas_totales * TOPE_ALTURA < total_pt:
            filas_totales += 1
        n_c = filas_totales - FILAS_COMENTARIO
        per_row_comentario = min(
            TOPE_ALTURA, max(ALTURA_BASE, ceil(total_pt / filas_totales))
        )

    if n_m > 0:
        _insertar_filas_materiales(ws, FILA_INSERCION_MATERIALES, n_m)
    if n_a > 0:
        _insertar_filas_actividades(
            ws,
            FILA_INSERCION_ACTIVIDADES + n_m,
            n_a,
            fila_estilo=FILA_ACTIVIDAD_INICIO + n_m,
        )
    if n_c > 0:
        _insertar_filas_comentarios(
            ws,
            FILA_INSERCION_COMENTARIOS + n_m + n_a,
            n_c,
            fila_estilo=FILA_COMENTARIO_INICIO + n_m + n_a,
        )
        _extender_merges_comentarios(
            ws, FILA_COMENTARIO_INICIO + n_m + n_a, FILA_COMENTARIO_INICIO + 2 + n_m + n_a + n_c
        )

    # Cabecera: FECHA como texto en la zona de negocio (D-8, REQ-REP-007).
    fecha_local = actividad.fecha.astimezone(zona)
    ws[CELDA_FECHA] = _limpiar_texto(
        f"{fecha_local.day:02d} / {_MESES[fecha_local.month - 1]} / {fecha_local.year}"
    )
    ws[CELDA_VEHICULO] = _limpiar_texto(actividad.vehiculo)
    ws[CELDA_NOMBRE] = _limpiar_texto(actividad.clientes)
    ws[CELDA_FOLIO] = _limpiar_texto(f"FOLIO: {actividad.folio}")
    ws[CELDA_DIRECCION] = _limpiar_texto(actividad.direcciones)

    # Materiales del snapshot: cantidades y valores como NÚMEROS, tal cual.
    for i, material in enumerate(materiales):
        fila = FILA_MATERIAL_INICIO + i
        ws.cell(row=fila, column=1, value=material.cantidad_consumida)
        if material.u_m is not None:
            ws.cell(row=fila, column=2, value=_limpiar_texto(material.u_m))
        celda_descripcion = ws.cell(
            row=fila, column=3, value=_limpiar_texto(material.descripcion)
        )
        if len(material.descripcion) > cpl_material:
            celda_descripcion.alignment = Alignment(wrap_text=True)
            ws.row_dimensions[fila].height = _alto_para_lineas(
                _lineas_texto(material.descripcion, cpl_material)
            )
        celda_valor = ws.cell(row=fila, column=9, value=material.valor_unitario)
        celda_valor.number_format = FORMATO_VALOR_UNITARIO

    # Actividades realizadas: una por fila, texto envuelto.
    for i, texto in enumerate(realizadas):
        fila = FILA_ACTIVIDAD_INICIO + n_m + i
        celda = ws.cell(row=fila, column=1, value=_limpiar_texto(texto))
        celda.alignment = Alignment(wrap_text=True, horizontal="left", vertical="top")
        ws.row_dimensions[fila].height = _alto_para_lineas(
            _lineas_texto(texto, cpl_actividad)
        )

    # Comentarios: valor íntegro en C32(+offsets), alturas repartidas.
    if comentarios:
        fila = FILA_COMENTARIO_INICIO + n_m + n_a
        celda = ws.cell(row=fila, column=3, value=_limpiar_texto(comentarios))
        celda.alignment = Alignment(wrap_text=True, horizontal="left", vertical="top")
        for fila_area in range(fila, fila + FILAS_COMENTARIO + n_c):
            ws.row_dimensions[fila_area].height = per_row_comentario


def generar_reporte_actividades(
    *, actividades: Sequence[ActividadOut], zona: ZoneInfo
) -> bytes:
    """Genera el libro XLSX del reporte de actividades inyectando los datos
    sobre la plantilla corporativa (una hoja por folio, orden folio
    ascendente, logo en todas las hojas) y devuelve los bytes del archivo.

    La plantilla en disco nunca se modifica; el resultado vive solo en
    memoria.
    """
    if not actividades:
        raise ValueError("No hay actividades para generar el reporte")
    ordenadas = sorted(actividades, key=lambda a: a.folio_num)

    wb = load_workbook(TEMPLATES_DIR / PLANTILLA_ACTIVIDADES)
    base = wb[HOJA_PLANTILLA]
    if HOJA_RESIDUAL in wb.sheetnames:
        wb.remove(wb[HOJA_RESIDUAL])

    # Bytes del logo capturados UNA sola vez, antes de copiar o guardar
    # (prohibido Image._data(): cierra el flujo de un solo uso).
    logo = cast(Any, base)._images[0]
    logo_bytes: bytes = logo.ref.getvalue()
    logo_ancho: int = logo.width
    logo_alto: int = logo.height

    # Todas las copias parten de la plantilla limpia (antes de inyectar).
    hojas: list[Worksheet] = [base] + [
        wb.copy_worksheet(base) for _ in range(len(ordenadas) - 1)
    ]
    for hoja, actividad in zip(hojas, ordenadas):
        hoja.title = actividad.folio
        if hoja is not base:
            _incrustar_logo(hoja, logo_bytes, logo_ancho, logo_alto)

    for hoja, actividad in zip(hojas, ordenadas):
        _llenar_hoja(hoja, actividad, zona)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def nombre_archivo_actividades(
    *, actividades: Sequence[ActividadOut], zona: ZoneInfo
) -> str:
    """Nombre del adjunto, construido solo con constantes, el folio y la
    fecha (C-5): `reporte-actividades-R0001.xlsx` con una hoja, o
    `reporte-actividades-AAAA-MM-DD.xlsx` (fecha de la actividad más
    reciente en la zona de negocio) con dos o más."""
    if not actividades:
        raise ValueError("No hay actividades para nombrar el reporte")
    if len(actividades) == 1:
        return sanitizar_nombre_archivo(
            f"reporte-actividades-{actividades[0].folio}.xlsx"
        )
    fecha_max = max(a.fecha for a in actividades).astimezone(zona).date().isoformat()
    return sanitizar_nombre_archivo(f"reporte-actividades-{fecha_max}.xlsx")
