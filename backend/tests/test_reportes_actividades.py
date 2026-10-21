"""Reporte de ACTIVIDADES (paquete 2026-10-actividades-reportes, Etapa 2/3)
— pruebas de API de `GET /api/v1/reportes/actividades` (T6, Fase 7).

Trazabilidad: cada prueba lleva su REQ-REP en nombre y docstring. Regla 8 de
tasks.md: cada prueba asevera la respuesta HTTP **y** el estado persistido
por SQL directo con `session` (incluidos los rechazos: «no se escribió
nada»). La siembra masiva o de borde se hace por INSERT directo (permitido
explícitamente en T6); los casos puntuales usan la API pública (helpers de
test_actividades). Sin mocks: todo contra el endpoint real.

Zona de negocio: `America/Mexico_City` (UTC-6 constante, C-1). Bordes de
día: 2026-10-20T05:30Z = 23:30 local del 19-oct; 2026-10-20T06:30Z =
00:30 local del 20-oct.
"""

import io
import json
import time
from datetime import date, datetime, timedelta, timezone
from typing import Any, cast
from zoneinfo import ZoneInfo

from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text

from tests.helpers.fabrica import (
    borrador_teams,
    carga_inicial,
    crear_equipo,
    crear_material,
    patch_material,
    procesar,
)
from tests.test_actividades import (
    _material_con_stock_en_equipo,
    _post_actividad,
    _seguimiento_cerrado,
)

UTC = timezone.utc
ZONA = ZoneInfo("America/Mexico_City")  # zona de negocio (C-1)
RUTA = "/api/v1/reportes/actividades"

# Nombres de mes del formato FECHA del XLSX (paridad con plantilla_reporte).
_MESES = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
    "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)  # fmt: skip

# INSERT de actividades para siembra masiva (CAST explícito para los jsonb).
SQL_INSERT_ACTIVIDAD = (
    "INSERT INTO actividades_reportes (folio_num, seguimiento_id, equipo_id, "
    "vehiculo, clientes, direcciones, actividades, comentarios, materiales, "
    "usuario, created_at) VALUES (:folio, :seg, :eq, 'Camioneta 12', "
    "'Cliente A', 'Calle 1', CAST(:act AS jsonb), NULL, CAST(:mat AS jsonb), "
    "'qa', :ca)"
)


def _entrada_material(i: int) -> dict:
    """Línea de material válida para el jsonb `materiales` (mismo esquema
    que `ActividadMaterialOut`)."""
    return {
        "material_id": i,
        "codigo": None,
        "descripcion": f"Material {i}",
        "u_m": "PZ",
        "cantidad_consumida": 1.0,
        "valor_unitario": 1.0,
    }


def _fecha_texto(created_at: datetime) -> str:
    """FECHA del XLSX: `DD / MES / AAAA` en la zona de negocio (REQ-REP-007)."""
    local = created_at.astimezone(ZONA)
    return f"{local.day:02d} / {_MESES[local.month - 1]} / {local.year}"


async def _contar(session, sql: str, **params) -> int:
    async with session.begin():
        return (await session.execute(text(sql), params)).scalar_one()


async def _conteo_bd(
    session,
    *,
    equipo_id: int | None = None,
    actividad_id: int | None = None,
    desde: date | None = None,
    hasta: date | None = None,
) -> int:
    """Conteo directo en BD con la MISMA semántica del servicio (días
    completos de la zona de negocio, filtros AND, `+ 1 day` en SQL)."""
    sql = "SELECT count(*) FROM actividades_reportes WHERE 1=1"
    params: dict[str, Any] = {}
    if equipo_id is not None:
        sql += " AND equipo_id = :equipo_id"
        params["equipo_id"] = equipo_id
    if actividad_id is not None:
        sql += " AND id = :actividad_id"
        params["actividad_id"] = actividad_id
    if desde is not None:
        sql += " AND created_at >= timezone(:zona, CAST(:desde AS timestamp))"
        params["desde"] = desde
    if hasta is not None:
        sql += (
            " AND created_at < timezone(:zona, "
            "CAST(:hasta AS timestamp) + interval '1 day')"
        )
        params["hasta"] = hasta
    if desde is not None or hasta is not None:
        params["zona"] = str(ZONA)
    return await _contar(session, sql, **params)


async def _insertar_seguimiento_sql(session, equipo_id: int) -> int:
    """Siembra DIRECTA por SQL de un seguimiento CERRADA (T6 permite INSERT
    directo para casos de borde/masivos)."""
    async with session.begin():
        return (
            await session.execute(
                text(
                    "INSERT INTO seguimientos (equipo_id, usuario, estado, "
                    "closed_at) VALUES (:e, 'qa', 'CERRADA', now()) "
                    "RETURNING id"
                ),
                {"e": equipo_id},
            )
        ).scalar_one()


async def _insertar_actividad_sql(
    session,
    *,
    equipo_id: int,
    folio_num: int,
    seguimiento_id: int,
    created_at: datetime,
    vehiculo: str = "Camioneta 12",
    clientes: str = "Cliente A",
    direcciones: str = "Calle 1",
    actividades: list[str] | None = None,
    comentarios: str | None = None,
    materiales: list[dict] | None = None,
    usuario: str = "qa",
) -> int:
    """Siembra DIRECTA por SQL de una actividad (created_at explícito)."""
    async with session.begin():
        return (
            await session.execute(
                text(
                    "INSERT INTO actividades_reportes (folio_num, "
                    "seguimiento_id, equipo_id, vehiculo, clientes, "
                    "direcciones, actividades, comentarios, materiales, "
                    "usuario, created_at) VALUES (:f, :s, :e, :v, :c, :d, "
                    "CAST(:act AS jsonb), :com, CAST(:mat AS jsonb), :u, :ca) "
                    "RETURNING id"
                ),
                {
                    "f": folio_num,
                    "s": seguimiento_id,
                    "e": equipo_id,
                    "v": vehiculo,
                    "c": clientes,
                    "d": direcciones,
                    "act": json.dumps(
                        actividades if actividades is not None
                        else ["Tendido de fibra"]
                    ),
                    "com": comentarios,
                    "mat": json.dumps(materiales if materiales is not None else []),
                    "u": usuario,
                    "ca": created_at,
                },
            )
        ).scalar_one()


async def _eventos_totales(session) -> int:
    return await _contar(session, "SELECT count(*) FROM auditoria_eventos")


# ---------------------------------------------------------------- REQ-REP-002


async def test_req_rep_002_json_lista_ordenada_y_total(
    client: AsyncClient, client_admin: AsyncClient, session
) -> None:
    """REQ-REP-002: json devuelve las actividades confirmadas ordenadas de la
    más reciente a la más antigua, con folio, folio_num, nombre de Unidad y
    materiales con valor unitario congelado; total == filas en BD."""
    material, equipo = await _material_con_stock_en_equipo(client)
    assert (
        await patch_material(client_admin, material["id_lista"], valor_unitario=7.5)
    ).status_code == 200
    seg1 = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=4
    )
    r1 = await _post_actividad(
        client, equipo["equipo_id"], seg1, vehiculo="Vehiculo Uno"
    )
    assert r1.status_code == 201, r1.text
    assert (
        await patch_material(client_admin, material["id_lista"], valor_unitario=9.999)
    ).status_code == 200
    seg2 = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=2
    )
    r2 = await _post_actividad(
        client, equipo["equipo_id"], seg2, vehiculo="Vehiculo Dos"
    )
    assert r2.status_code == 201, r2.text

    resp = await client.get(RUTA, params={"formato": "json"})
    assert resp.status_code == 200, resp.text
    cuerpo = resp.json()
    assert cuerpo["total"] == 2
    # Más reciente primero (la segunda confirmación encabeza).
    assert [i["actividad_id"] for i in cuerpo["items"]] == [
        r2.json()["actividad_id"], r1.json()["actividad_id"],
    ]
    primero = cuerpo["items"][0]
    assert primero["folio"] == "R0002"
    assert primero["folio_num"] == 2
    assert primero["equipo_nombre"] == "Unidad Act"
    assert primero["vehiculo"] == "Vehiculo Dos"
    # Valor unitario congelado del snapshot (REQ-ACT-010 traslado a reporte).
    assert primero["materiales"][0]["valor_unitario"] == 9.999
    assert cuerpo["items"][1]["materiales"][0]["valor_unitario"] == 7.5

    # BD: las dos filas con el mismo orden de la consulta del servicio.
    async with session.begin():
        filas = (
            await session.execute(
                text(
                    "SELECT folio_num, created_at FROM actividades_reportes "
                    "ORDER BY created_at DESC, id DESC"
                )
            )
        ).all()
    assert [f.folio_num for f in filas] == [2, 1]
    assert filas[0].created_at >= filas[1].created_at
    assert await _contar(
        session, "SELECT count(*) FROM actividades_reportes"
    ) == 2


async def test_req_rep_002_filtros_and(client: AsyncClient, session) -> None:
    """REQ-REP-002: los filtros se combinan con semántica AND —equipo_id +
    desde + hasta, equipo_id + actividad_id— y una actividad que no cumple
    otro filtro queda fuera (cada consulta se contrasta con SQL directo)."""
    material, equipo_a = await _material_con_stock_en_equipo(client)
    equipo_b = await crear_equipo(client, nombre="Unidad B Filtros")
    b = await borrador_teams(
        client,
        almacen_id=1,
        equipo_id=equipo_b["equipo_id"],
        material_id=material["id_lista"],
        cantidad=20,
    )
    await procesar(client, b["id"])

    s_a1 = await _seguimiento_cerrado(
        client, equipo_a["equipo_id"], material["id_lista"], consumida=1
    )
    a1 = await _post_actividad(client, equipo_a["equipo_id"], s_a1)
    s_b1 = await _seguimiento_cerrado(
        client, equipo_b["equipo_id"], material["id_lista"], consumida=1
    )
    b1 = await _post_actividad(client, equipo_b["equipo_id"], s_b1)
    s_a2 = await _seguimiento_cerrado(
        client, equipo_a["equipo_id"], material["id_lista"], consumida=1
    )
    a2 = await _post_actividad(client, equipo_a["equipo_id"], s_a2)
    assert a1.status_code == b1.status_code == a2.status_code == 201

    hoy = datetime.now(UTC).astimezone(ZONA).date()
    manana = hoy + timedelta(days=1)
    id_a1, id_a2, id_b1 = (
        a1.json()["actividad_id"],
        a2.json()["actividad_id"],
        b1.json()["actividad_id"],
    )

    # equipo_id + desde + hasta: solo las dos de la Unidad A.
    resp = await client.get(
        RUTA,
        params={
            "equipo_id": equipo_a["equipo_id"],
            "desde": hoy.isoformat(),
            "hasta": hoy.isoformat(),
            "formato": "json",
        },
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 2
    assert {i["actividad_id"] for i in resp.json()["items"]} == {id_a1, id_a2}
    assert (
        await _conteo_bd(session, equipo_id=equipo_a["equipo_id"], desde=hoy, hasta=hoy)
        == 2
    )

    # equipo_id + actividad_id: solo esa actividad.
    resp = await client.get(
        RUTA,
        params={
            "equipo_id": equipo_a["equipo_id"],
            "actividad_id": id_a2,
            "formato": "json",
        },
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 1
    assert resp.json()["items"][0]["actividad_id"] == id_a2
    assert (
        await _conteo_bd(session, equipo_id=equipo_a["equipo_id"], actividad_id=id_a2)
        == 1
    )

    # actividad_id que NO cumple el rango: vacío.
    resp = await client.get(
        RUTA, params={"actividad_id": id_a2, "desde": manana.isoformat(), "formato": "json"}
    )
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}
    assert await _conteo_bd(session, actividad_id=id_a2, desde=manana) == 0

    # actividad_id de otra Unidad combinado con equipo_id: vacío.
    resp = await client.get(
        RUTA,
        params={
            "equipo_id": equipo_a["equipo_id"],
            "actividad_id": id_b1,
            "formato": "json",
        },
    )
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}
    assert (
        await _conteo_bd(session, equipo_id=equipo_a["equipo_id"], actividad_id=id_b1)
        == 0
    )


async def test_req_rep_002_json_404_no_aplica(client: AsyncClient, session) -> None:
    """REQ-REP-002/010: json sin resultados responde 200 {items: [], total: 0}
    (nunca 404); xlsx sin resultados responde 404 canónico. BD sin cambios."""
    equipo = await crear_equipo(client, nombre="Unidad Vacia")
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 0

    resp = await client.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "json"}
    )
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"items": [], "total": 0}

    resp = await client.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "xlsx"}
    )
    assert resp.status_code == 404, resp.text
    error = resp.json()["error"]
    assert set(error) == {"code", "message", "coordinates"}
    assert error["coordinates"] == [{"formato": "xlsx"}]

    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 0
    assert await _eventos_totales(session) == 0


# ---------------------------------------------------------------- REQ-REP-003


async def test_req_rep_003_actor_sin_acceso_403(
    client: AsyncClient, client_sin_scope: AsyncClient, session
) -> None:
    """REQ-REP-003: Unidad ajena ⇒ 403 con cuerpo de autorización; sin
    equipo_id ⇒ 200 vacío (actor sin Unidades); actividad_id ajeno ⇒
    inexistente (vacío/404). BD intacta."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    act = await _post_actividad(client, equipo["equipo_id"], seg)
    assert act.status_code == 201
    actividad_id = act.json()["actividad_id"]
    eventos_antes = await _eventos_totales(session)

    resp = await client_sin_scope.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "json"}
    )
    assert resp.status_code == 403, resp.text
    error = resp.json()["error"]
    assert error["code"] == "AUTHORIZATION_FAILED"
    assert error["actor"] == "intruso"

    # Sin equipo_id: actor sin Unidades visibles ⇒ 200 vacío.
    resp = await client_sin_scope.get(RUTA, params={"formato": "json"})
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"items": [], "total": 0}

    # actividad_id de Unidad ajena: se trata como inexistente, sin filtrar.
    resp = await client_sin_scope.get(
        RUTA, params={"actividad_id": actividad_id, "formato": "json"}
    )
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}
    resp = await client_sin_scope.get(
        RUTA, params={"actividad_id": actividad_id, "formato": "xlsx"}
    )
    assert resp.status_code == 404, resp.text

    # BD: nada nuevo (una actividad, sin eventos de autorización fallida).
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1
    assert await _eventos_totales(session) == eventos_antes


async def test_req_rep_003_admin_transversal_y_equipo_inexistente(
    client: AsyncClient,
    client_admin: AsyncClient,
    client_sin_scope: AsyncClient,
    session,
) -> None:
    """REQ-REP-003: el administrador ve Unidades ajenas; con equipo_id
    inexistente recibe vacío (json) / 404 (xlsx); el actor no administrador
    recibe 403 indistinguible del de una Unidad ajena."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    assert (await _post_actividad(client, equipo["equipo_id"], seg)).status_code == 201

    # Administrador sin membresía: ve los datos de la Unidad.
    resp = await client_admin.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "json"}
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 1
    assert resp.json()["items"][0]["equipo_nombre"] == "Unidad Act"

    # Administrador + equipo_id inexistente: vacío en json, 404 en xlsx.
    resp = await client_admin.get(RUTA, params={"equipo_id": 999999, "formato": "json"})
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}
    resp = await client_admin.get(RUTA, params={"equipo_id": 999999, "formato": "xlsx"})
    assert resp.status_code == 404, resp.text
    assert resp.json()["error"]["coordinates"] == [{"formato": "xlsx"}]

    # No administrador + equipo_id inexistente: 403 (sin prechequeo 404).
    resp = await client_sin_scope.get(
        RUTA, params={"equipo_id": 999999, "formato": "json"}
    )
    assert resp.status_code == 403, resp.text

    # BD sin cambios.
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


async def test_req_rep_003_unidad_inactiva_incluida(
    client: AsyncClient, client_admin: AsyncClient, session
) -> None:
    """REQ-REP-003 (C-7): una actividad de una Unidad dada de baja sigue en
    el reporte (el historial documental se conserva; sin prechequeo de
    Unidad activa)."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    act = await _post_actividad(client, equipo["equipo_id"], seg)
    assert act.status_code == 201

    # Baja de la Unidad (soft delete: is_active = false).
    resp = await client.delete(f"/api/v1/equipos/{equipo['equipo_id']}")
    assert resp.status_code == 204, resp.text
    async with session.begin():
        activa = (
            await session.execute(
                text("SELECT is_active FROM equipos WHERE equipo_id = :e"),
                {"e": equipo["equipo_id"]},
            )
        ).scalar_one()
    assert activa is False

    resp = await client_admin.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "json"}
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 1
    assert resp.json()["items"][0]["actividad_id"] == act.json()["actividad_id"]
    assert await _conteo_bd(session, equipo_id=equipo["equipo_id"]) == 1


# ---------------------------------------------------------------- REQ-REP-004


async def test_req_rep_004_dia_inclusivo_y_bordes_zona(
    client: AsyncClient, session
) -> None:
    """REQ-REP-004: desde/hasta son días calendario completos de la zona de
    negocio (UTC-6 constante) con ambos extremos inclusivos e independientes.
    Borde: 05:30Z = 23:30 local del 19-oct; 06:30Z = 00:30 local del 20-oct.
    Cubre también el caso de solo-desde y solo-hasta."""
    equipo = await crear_equipo(client, nombre="Unidad Bordes")
    eid = equipo["equipo_id"]
    fecha_19 = datetime(2026, 10, 20, 5, 30, tzinfo=UTC)  # 23:30 local 19-oct
    fecha_20 = datetime(2026, 10, 20, 6, 30, tzinfo=UTC)  # 00:30 local 20-oct
    seg19 = await _insertar_seguimiento_sql(session, eid)
    seg20 = await _insertar_seguimiento_sql(session, eid)
    await _insertar_actividad_sql(
        session, equipo_id=eid, folio_num=1, seguimiento_id=seg19,
        created_at=fecha_19, vehiculo="Borde 19",
    )
    await _insertar_actividad_sql(
        session, equipo_id=eid, folio_num=2, seguimiento_id=seg20,
        created_at=fecha_20, vehiculo="Borde 20",
    )
    d19, d20 = date(2026, 10, 19), date(2026, 10, 20)

    async def consultar(**params) -> dict:
        resp = await client.get(RUTA, params={"formato": "json", **params})
        assert resp.status_code == 200, resp.text
        return resp.json()

    # desde=hasta=19: entra la de 23:30 del 19; la de 00:30 del 20 no.
    cuerpo = await consultar(desde="2026-10-19", hasta="2026-10-19")
    assert [i["vehiculo"] for i in cuerpo["items"]] == ["Borde 19"]
    assert await _conteo_bd(session, desde=d19, hasta=d19) == 1

    # desde=19 + hasta=20: entran ambas (00:30 del 20 < inicio del 21).
    cuerpo = await consultar(desde="2026-10-19", hasta="2026-10-20")
    assert cuerpo["total"] == 2
    assert [i["vehiculo"] for i in cuerpo["items"]] == ["Borde 20", "Borde 19"]
    assert await _conteo_bd(session, desde=d19, hasta=d20) == 2

    # solo-desde=19: ambas (>= inicio del 19, sin cota superior).
    cuerpo = await consultar(desde="2026-10-19")
    assert cuerpo["total"] == 2
    assert await _conteo_bd(session, desde=d19) == 2

    # solo-hasta=19: solo la del 19 (< inicio del 20).
    cuerpo = await consultar(hasta="2026-10-19")
    assert cuerpo["total"] == 1
    assert [i["vehiculo"] for i in cuerpo["items"]] == ["Borde 19"]
    assert await _conteo_bd(session, hasta=d19) == 1

    # solo-desde=20: solo la de 00:30 del 20 (la de 23:30 del 19 queda fuera).
    cuerpo = await consultar(desde="2026-10-20")
    assert [i["vehiculo"] for i in cuerpo["items"]] == ["Borde 20"]
    assert await _conteo_bd(session, desde=d20) == 1

    # solo-hasta=20: ambas (00:30 del 20 < inicio del 21).
    cuerpo = await consultar(hasta="2026-10-20")
    assert cuerpo["total"] == 2
    assert await _conteo_bd(session, hasta=d20) == 2


async def test_req_rep_004_extremos_de_fecha_sin_500(
    client: AsyncClient, session
) -> None:
    """REQ-REP-004 (F-04): desde=0001-01-01 y hasta=9999-12-31 responden 200
    (el «+ 1 día» se calcula en SQL, sin desbordar en Python)."""
    equipo = await crear_equipo(client, nombre="Unidad Extremos")
    seg = await _insertar_seguimiento_sql(session, equipo["equipo_id"])
    await _insertar_actividad_sql(
        session, equipo_id=equipo["equipo_id"], folio_num=1,
        seguimiento_id=seg, created_at=datetime(2026, 10, 20, 12, 0, tzinfo=UTC),
    )
    d_min, d_max = date(1, 1, 1), date(9999, 12, 31)

    for params, esperado in (
        ({"desde": "0001-01-01", "hasta": "9999-12-31"}, 1),
        ({"desde": "0001-01-01"}, 1),
        ({"hasta": "9999-12-31"}, 1),
    ):
        resp = await client.get(RUTA, params={**params, "formato": "json"})
        assert resp.status_code == 200, (params, resp.text)
        assert resp.json()["total"] == esperado
    assert await _conteo_bd(session, desde=d_min, hasta=d_max) == 1
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


async def test_req_rep_004_desde_mayor_que_hasta_422(
    client: AsyncClient, client_sin_scope: AsyncClient, session
) -> None:
    """REQ-REP-004/011: desde > hasta ⇒ 422 RANGO_FECHAS_INVALIDO con cuerpo
    canónico; la validación precede a la autorización (mismo 422 con Unidad
    ajena) y no consulta ni escribe datos."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    assert (await _post_actividad(client, equipo["equipo_id"], seg)).status_code == 201

    params = {"desde": "2026-10-20", "hasta": "2026-10-19", "formato": "json"}
    resp = await client.get(RUTA, params=params)
    assert resp.status_code == 422, resp.text
    error = resp.json()["error"]
    assert set(error) == {"code", "message", "coordinates"}
    assert error["code"] == "RANGO_FECHAS_INVALIDO"

    # Precedencia: validación ANTES de autorización ⇒ sigue 422 con Unidad ajena.
    resp = await client_sin_scope.get(
        RUTA, params={**params, "equipo_id": equipo["equipo_id"]}
    )
    assert resp.status_code == 422, resp.text
    assert resp.json()["error"]["code"] == "RANGO_FECHAS_INVALIDO"

    # BD sin cambios: la actividad creada sigue siendo la única fila.
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


# ---------------------------------------------------------------- REQ-REP-005


async def test_req_rep_005_por_operacion(client: AsyncClient, session) -> None:
    """REQ-REP-005: actividad_id aislado devuelve únicamente esa actividad
    (folio como etiqueta de la operación); combinado con un rango que la
    excluye produce vacío; BD con la fila exacta."""
    material, equipo = await _material_con_stock_en_equipo(client)
    s1 = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    a1 = await _post_actividad(
        client, equipo["equipo_id"], s1, vehiculo="Vehiculo Operacion 1"
    )
    s2 = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    a2 = await _post_actividad(
        client, equipo["equipo_id"], s2, vehiculo="Vehiculo Operacion 2"
    )
    assert a1.status_code == 201 and a2.status_code == 201
    id1 = a1.json()["actividad_id"]

    resp = await client.get(RUTA, params={"actividad_id": id1, "formato": "json"})
    assert resp.status_code == 200, resp.text
    cuerpo = resp.json()
    assert cuerpo["total"] == 1
    item = cuerpo["items"][0]
    assert item["actividad_id"] == id1
    assert item["folio"] == "R0001"  # etiqueta de la operación
    assert item["vehiculo"] == "Vehiculo Operacion 1"

    # AND con un rango que la excluye: vacío.
    manana = (datetime.now(UTC).astimezone(ZONA).date() + timedelta(days=1)).isoformat()
    resp = await client.get(
        RUTA, params={"actividad_id": id1, "desde": manana, "formato": "json"}
    )
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}

    # BD: la fila exacta y nada más.
    async with session.begin():
        fila = (
            await session.execute(
                text("SELECT folio_num, vehiculo FROM actividades_reportes WHERE id = :i"),
                {"i": id1},
            )
        ).one()
    assert fila.folio_num == 1
    assert fila.vehiculo == "Vehiculo Operacion 1"
    assert await _conteo_bd(session, actividad_id=id1) == 1


# ---------------------------------------------------------------- REQ-REP-008


async def test_req_rep_008_desborde_via_api(client: AsyncClient, session) -> None:
    """REQ-REP-008: desborde real vía API (11 materiales, 5 actividades de
    500 caracteres y comentario de ~2000 con saltos) ⇒ xlsx 200 con las
    filas insertadas, celdas combinadas y texto íntegro; BD intacta tras la
    descarga."""
    materiales = [
        await crear_material(client, descripcion=f"Material Desborde {i:02d}")
        for i in range(1, 12)
    ]
    equipo = await crear_equipo(client, nombre="Unidad Desborde")
    eid = equipo["equipo_id"]
    for m in materiales:
        await carga_inicial(client, material_id=m["id_lista"], cantidad=100)
    mov = await client.post(
        "/api/v1/movimientos",
        json={
            "tipo_movimiento": "TEAMS",
            "origen_almacen_id": 1,
            "destino_equipo_id": eid,
            "detalle": [{"material_id": m["id_lista"], "cantidad": 50} for m in materiales],
        },
    )
    assert mov.status_code == 201, mov.text
    await procesar(client, mov.json()["id"])

    resp = await client.post(
        f"/api/v1/equipos/{eid}/seguimientos",
        json={
            "items": [
                {"material_id": m["id_lista"], "cantidad_tomada": 10}
                for m in materiales
            ]
        },
    )
    assert resp.status_code == 201, resp.text
    seg = resp.json()["id"]
    resp = await client.post(
        f"/api/v1/equipos/{eid}/seguimientos/{seg}/cerrar",
        json={
            "consumidos": [
                {"material_id": m["id_lista"], "cantidad_consumida": 5}
                for m in materiales
            ]
        },
    )
    assert resp.status_code == 200, resp.text

    celdas = [f"Actividad larga {i + 1}: " + "A" * 481 for i in range(5)]
    assert all(len(c) == 500 for c in celdas)
    comentario = "\n".join("P" * 38 for _ in range(50))  # 1949 chars con saltos
    assert len(comentario) == 1949
    resp = await _post_actividad(
        client, eid, seg, actividades=celdas, comentarios=comentario
    )
    assert resp.status_code == 201, resp.text

    tablas = (
        "actividades_reportes", "seguimientos", "seguimiento_items",
        "actas_reportes", "inventario_equipos", "auditoria_eventos",
    )
    antes = {t: await _contar(session, f"SELECT count(*) FROM {t}") for t in tablas}
    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "xlsx"})
    assert resp.status_code == 200, resp.text
    despues = {t: await _contar(session, f"SELECT count(*) FROM {t}") for t in tablas}
    assert despues == antes  # la descarga no escribe nada

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["R0001"]
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    # 11 filas de materiales (13..23) con C:H combinada y cantidad numérica.
    for fila in range(13, 24):
        assert f"C{fila}:H{fila}" in merges
        assert ws.cell(row=fila, column=1).value == 5
    # 5 filas de actividades (26..30) con A:I combinada y texto íntegro.
    for fila in range(26, 31):
        assert f"A{fila}:I{fila}" in merges
        assert ws.cell(row=fila, column=1).value == celdas[fila - 26]
    # Comentario íntegro en el área desplazada (32 + 1 material + 1 actividad).
    assert ws["C34"].value == comentario
    assert len(cast(Any, ws)._images) == 1  # logo conservado


# ---------------------------------------------------------------- REQ-REP-009


async def test_req_rep_009_xlsx_descargable_y_hojas(
    client: AsyncClient, client_admin: AsyncClient, session
) -> None:
    """REQ-REP-009/007: xlsx 200 con media type XLSX, adjunto con nombre
    `reporte-actividades-R0001.xlsx`, una hoja titulada con el folio,
    cabecera y materiales del snapshot (valor unitario congelado) y logo."""
    m1 = await crear_material(client, descripcion="Material A Xlsx")
    m2 = await crear_material(client, descripcion="Material B Xlsx")
    equipo = await crear_equipo(client, nombre="Unidad Xlsx")
    eid = equipo["equipo_id"]
    await carga_inicial(client, material_id=m1["id_lista"], cantidad=100)
    await carga_inicial(client, material_id=m2["id_lista"], cantidad=100)
    mov = await client.post(
        "/api/v1/movimientos",
        json={
            "tipo_movimiento": "TEAMS",
            "origen_almacen_id": 1,
            "destino_equipo_id": eid,
            "detalle": [
                {"material_id": m1["id_lista"], "cantidad": 50},
                {"material_id": m2["id_lista"], "cantidad": 50},
            ],
        },
    )
    assert mov.status_code == 201, mov.text
    await procesar(client, mov.json()["id"])
    assert (
        await patch_material(client_admin, m1["id_lista"], valor_unitario=12.34)
    ).status_code == 200
    assert (
        await patch_material(client_admin, m2["id_lista"], valor_unitario=0.5)
    ).status_code == 200

    resp = await client.post(
        f"/api/v1/equipos/{eid}/seguimientos",
        json={
            "items": [
                {"material_id": m1["id_lista"], "cantidad_tomada": 10},
                {"material_id": m2["id_lista"], "cantidad_tomada": 10},
            ]
        },
    )
    assert resp.status_code == 201, resp.text
    seg = resp.json()["id"]
    resp = await client.post(
        f"/api/v1/equipos/{eid}/seguimientos/{seg}/cerrar",
        json={
            "consumidos": [
                {"material_id": m1["id_lista"], "cantidad_consumida": 2.5},
                {"material_id": m2["id_lista"], "cantidad_consumida": 7},
            ]
        },
    )
    assert resp.status_code == 200, resp.text
    act = await _post_actividad(
        client, eid, seg,
        vehiculo="Camioneta 99",
        clientes="Cliente XLSX",
        direcciones="Av. XLSX 5",
        actividades=["Tarea 1", "Tarea 2"],
    )
    assert act.status_code == 201, act.text
    actividad_id = act.json()["actividad_id"]

    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "xlsx"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        resp.headers["content-disposition"]
        == 'attachment; filename="reporte-actividades-R0001.xlsx"'
    )

    # FECHA esperada desde el created_at persistido (zona de negocio).
    async with session.begin():
        creado = (
            await session.execute(
                text("SELECT created_at FROM actividades_reportes WHERE id = :i"),
                {"i": actividad_id},
            )
        ).scalar_one()

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["R0001"]
    assert "Sheet1" not in wb.sheetnames
    ws = wb["R0001"]
    assert ws["B8"].value == _fecha_texto(creado)
    assert ws["E8"].value == "Camioneta 99"
    assert ws["B9"].value == "Cliente XLSX"
    assert ws["I9"].value == "FOLIO: R0001"
    assert ws["B10"].value == "Av. XLSX 5"
    # Materiales del snapshot (números y valor unitario congelado).
    assert ws["A13"].value == 2.5
    assert ws["B13"].value == "PZ"
    assert ws["C13"].value == "Material A Xlsx"
    assert ws["I13"].value == 12.34
    assert ws["I13"].number_format == "#,##0.00#"
    assert ws["A14"].value == 7
    assert ws["C14"].value == "Material B Xlsx"
    assert ws["I14"].value == 0.5
    assert ws["A25"].value == "Tarea 1"
    assert ws["A26"].value == "Tarea 2"
    assert len(cast(Any, ws)._images) == 1  # logo corporativo

    # BD: la fila exacta del reporte.
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1
    assert await _conteo_bd(session, equipo_id=eid) == 1


async def test_req_rep_009_101_folios_422(client: AsyncClient, session) -> None:
    """REQ-REP-009: más de 100 folios en xlsx ⇒ 422 LIMITE_HOJAS_EXCEDIDO con
    coordinates y sin generar libro; el json lista las 101 en orden de fecha
    descendente; BD con 101 filas."""
    equipo = await crear_equipo(client, nombre="Unidad Masiva")
    eid = equipo["equipo_id"]
    base = datetime(2026, 10, 1, 12, 0, tzinfo=UTC)
    async with session.begin():
        seguimientos = (
            await session.execute(
                text(
                    "INSERT INTO seguimientos (equipo_id, usuario, estado, "
                    "closed_at) SELECT :e, 'qa', 'CERRADA', now() "
                    "FROM generate_series(1, 101) RETURNING id"
                ),
                {"e": eid},
            )
        ).scalars().all()
        filas = [
            {
                "folio": i,
                "seg": sid,
                "eq": eid,
                "act": json.dumps([f"Actividad del folio {i}"]),
                "mat": json.dumps([_entrada_material(i)]),
                "ca": base + timedelta(minutes=i),
            }
            for i, sid in enumerate(seguimientos, start=1)
        ]
        await session.execute(text(SQL_INSERT_ACTIVIDAD), filas)

    # json: sin tope (C-6), las 101 en orden de fecha descendente.
    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "json"})
    assert resp.status_code == 200, resp.text
    cuerpo = resp.json()
    assert cuerpo["total"] == 101
    assert [i["folio_num"] for i in cuerpo["items"]] == list(range(101, 0, -1))
    assert cuerpo["items"][0]["folio"] == "R0101"

    # xlsx: 422 canónico, sin archivo.
    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "xlsx"})
    assert resp.status_code == 422, resp.text
    error = resp.json()["error"]
    assert set(error) == {"code", "message", "coordinates"}
    assert error["code"] == "LIMITE_HOJAS_EXCEDIDO"
    assert {"formato": "xlsx", "actividades": 101} in error["coordinates"]
    assert resp.headers["content-type"].startswith("application/json")
    assert not resp.content.startswith(b"PK")  # sin libro generado

    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 101


async def test_req_rep_009_100_hojas_tiempo_y_tamano(
    client: AsyncClient, session
) -> None:
    """REQ-REP-009/NF: libro de 100 hojas ⇒ 200, >100 KB, abre con openpyxl
    con 100 hojas tituladas R0001..R0100 y una imagen de logo por hoja. Se
    mide tiempo y tamaño (evidencia T6, sin asertos frágiles de tiempo)."""
    equipo = await crear_equipo(client, nombre="Unidad Cien")
    eid = equipo["equipo_id"]
    base = datetime(2026, 10, 5, 12, 0, tzinfo=UTC)
    async with session.begin():
        seguimientos = (
            await session.execute(
                text(
                    "INSERT INTO seguimientos (equipo_id, usuario, estado, "
                    "closed_at) SELECT :e, 'qa', 'CERRADA', now() "
                    "FROM generate_series(1, 100) RETURNING id"
                ),
                {"e": eid},
            )
        ).scalars().all()
        filas = [
            {
                "folio": i,
                "seg": sid,
                "eq": eid,
                "act": json.dumps([f"Actividad del folio {i}"]),
                "mat": json.dumps([_entrada_material(i)]),
                "ca": base + timedelta(minutes=i),
            }
            for i, sid in enumerate(seguimientos, start=1)
        ]
        await session.execute(text(SQL_INSERT_ACTIVIDAD), filas)

    inicio = time.perf_counter()
    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "xlsx"})
    duracion = time.perf_counter() - inicio
    assert resp.status_code == 200, resp.text
    tamano = len(resp.content)
    print(f"[T6] libro 100 hojas: tiempo={duracion:.2f}s tamano={tamano:,} bytes")
    assert tamano > 100_000

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == [f"R{i:04d}" for i in range(1, 101)]
    assert "Sheet1" not in wb.sheetnames
    for nombre in wb.sheetnames:
        assert len(cast(Any, wb[nombre])._images) == 1
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 100


# ---------------------------------------------------------------- REQ-REP-010


async def test_req_rep_010_limite_filas_422(client: AsyncClient, session) -> None:
    """REQ-REP-010: una actividad con 10 001 materiales (10 002 filas
    inyectadas) ⇒ xlsx 422 LIMITE_FILAS_EXCEDIDO; el json sigue devolviendo
    la actividad (el tope aplica solo al libro). Una sola hoja: no se rompe
    el tope de hojas."""
    equipo = await crear_equipo(client, nombre="Unidad Filas")
    eid = equipo["equipo_id"]
    seg = await _insertar_seguimiento_sql(session, eid)
    await _insertar_actividad_sql(
        session,
        equipo_id=eid,
        folio_num=1,
        seguimiento_id=seg,
        created_at=datetime(2026, 10, 20, 12, 0, tzinfo=UTC),
        materiales=[_entrada_material(i) for i in range(1, 10002)],
    )

    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "xlsx"})
    assert resp.status_code == 422, resp.text
    error = resp.json()["error"]
    assert set(error) == {"code", "message", "coordinates"}
    assert error["code"] == "LIMITE_FILAS_EXCEDIDO"
    assert any(
        c.get("formato") == "xlsx" and c.get("filas") == 10002
        for c in error["coordinates"]
    )
    assert not resp.content.startswith(b"PK")

    resp = await client.get(RUTA, params={"equipo_id": eid, "formato": "json"})
    assert resp.status_code == 200, resp.text
    cuerpo = resp.json()
    assert cuerpo["total"] == 1
    assert len(cuerpo["items"][0]["materiales"]) == 10001

    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


async def test_req_rep_010_formulas_y_nombre_en_descarga(
    client: AsyncClient, session
) -> None:
    """REQ-REP-010 (C-10): los textos que inician con =/+/-/@ se neutralizan
    con apóstrofe literal en el XLSX (cero celdas fórmula en todo el libro)
    y el adjunto se nombra con el folio."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=1
    )
    resp = await _post_actividad(
        client, equipo["equipo_id"], seg,
        vehiculo="=cmd()",
        clientes="- Instalación",
        direcciones="@calle",
        comentarios="+suma",
    )
    assert resp.status_code == 201, resp.text

    resp = await client.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "xlsx"}
    )
    assert resp.status_code == 200, resp.text
    assert (
        resp.headers["content-disposition"]
        == 'attachment; filename="reporte-actividades-R0001.xlsx"'
    )

    wb = load_workbook(io.BytesIO(resp.content))
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                assert celda.data_type != "f"
    ws = wb["R0001"]
    assert ws["E8"].value == "'=cmd()"
    assert ws["B9"].value == "'- Instalación"
    assert ws["B10"].value == "'@calle"
    assert ws["C32"].value == "'+suma"
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


# ---------------------------------------------------------------- REQ-REP-011


async def test_req_rep_011_validaciones_422(client: AsyncClient, session) -> None:
    """REQ-REP-011: formato inválido (literal_error), parámetro no declarado
    (extra_forbidden), identificadores <= 0 (greater_than), fecha inválida y
    X-Actor ausente responden 422 sin tocar datos; con actor válido y sin
    datos, 200 vacío."""
    casos: list[tuple[dict, str]] = [
        ({"formato": "zip"}, "literal_error"),
        ({"acta_id": 1}, "extra_forbidden"),
        ({"equipo_id": 0}, "greater_than"),
        ({"actividad_id": 0}, "greater_than"),
    ]
    for params, tipo in casos:
        resp = await client.get(RUTA, params=params)
        assert resp.status_code == 422, (params, resp.text)
        detalle = resp.json()["detail"]
        assert any(e["type"] == tipo for e in detalle), (params, detalle)

    resp = await client.get(RUTA, params={"desde": "2026-99-99"})
    assert resp.status_code == 422, resp.text

    # X-Actor ausente: 422 de validación del header, sin consultar BD.
    from app.main import app

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as sin_header:
        resp = await sin_header.get(RUTA, params={"formato": "json"})
    assert resp.status_code == 422, resp.text
    detalle = resp.json()["detail"]
    assert any(
        e["loc"][0] == "header" and str(e["loc"][1]).lower() == "x-actor"
        for e in detalle
    ), detalle

    # BD: cero filas y cero eventos (ninguno de los rechazos escribió nada).
    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 0
    assert await _eventos_totales(session) == 0

    # Con actor válido y sin datos: 200 vacío (sin Unidades visibles).
    resp = await client.get(RUTA, params={"formato": "json"})
    assert resp.status_code == 200
    assert resp.json() == {"items": [], "total": 0}


async def test_req_rep_011_identificadores_fuera_de_int4_400(
    client_admin: AsyncClient, session
) -> None:
    """REQ-REP-011 (F-04): identificadores positivos fuera de INT4 (10^12,
    10^30) responden el 400 canónico NUMERIC_OUT_OF_RANGE (SQLSTATE 22003),
    nunca 500, en json y xlsx; BD sin cambios."""
    equipo = await crear_equipo(client_admin, nombre="Unidad Int4")
    seg = await _insertar_seguimiento_sql(session, equipo["equipo_id"])
    await _insertar_actividad_sql(
        session, equipo_id=equipo["equipo_id"], folio_num=1,
        seguimiento_id=seg, created_at=datetime(2026, 10, 20, 12, 0, tzinfo=UTC),
    )

    for param, valor in (("equipo_id", 10**12), ("actividad_id", 10**30)):
        for formato in ("json", "xlsx"):
            resp = await client_admin.get(
                RUTA, params={param: valor, "formato": formato}
            )
            assert resp.status_code == 400, (param, formato, resp.text)
            error = resp.json()["error"]
            assert set(error) == {"code", "message", "coordinates"}
            assert error["code"] == "NUMERIC_OUT_OF_RANGE"

    assert await _contar(session, "SELECT count(*) FROM actividades_reportes") == 1


# ---------------------------------------------------------------- REQ-REP-016

TABLAS_SNAPSHOT = [
    "actividades_reportes",
    "inventario_equipos",
    "seguimientos",
    "seguimiento_items",
    "actas_reportes",
    "auditoria_eventos",
]

_ORDEN_SNAPSHOT = {
    "actividades_reportes": "id",
    "inventario_equipos": "equipo_id, material_id",
    "seguimientos": "id",
    "seguimiento_items": "id",
    "actas_reportes": "id",
    "auditoria_eventos": "id",
}


async def _snapshot(session) -> dict[str, list[tuple]]:
    """Volcado íntegro y ordenado de las tablas vigiladas por REQ-REP-016."""
    async with session.begin():
        snap: dict[str, list[tuple]] = {}
        for tabla in TABLAS_SNAPSHOT:
            filas = (
                await session.execute(
                    text(f"SELECT * FROM {tabla} ORDER BY {_ORDEN_SNAPSHOT[tabla]}")
                )
            ).all()
            snap[tabla] = [tuple(f) for f in filas]
    return snap


async def test_req_rep_016_solo_lectura(client: AsyncClient, session) -> None:
    """REQ-REP-016: GET json y GET xlsx no modifican actividades_reportes,
    inventario_equipos, seguimientos, seguimiento_items, actas_reportes ni
    auditoria_eventos (snapshots idénticos, cero eventos nuevos)."""
    material, equipo = await _material_con_stock_en_equipo(client)
    seg = await _seguimiento_cerrado(
        client, equipo["equipo_id"], material["id_lista"], consumida=2
    )
    assert (await _post_actividad(client, equipo["equipo_id"], seg)).status_code == 201

    antes = await _snapshot(session)

    resp = await client.get(RUTA, params={"formato": "json"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 1
    resp = await client.get(
        RUTA, params={"equipo_id": equipo["equipo_id"], "formato": "xlsx"}
    )
    assert resp.status_code == 200, resp.text

    despues = await _snapshot(session)
    assert despues == antes
