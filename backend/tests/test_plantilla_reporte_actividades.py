"""Servicio PlantillaReporteActividades — REQ-REP-006/007/008/009/010.

Verifica que el XLSX de Actividades se obtiene INYECTANDO datos sobre la
plantilla corporativa «REPORTE DE ACTIVIDADES.xlsx» (jamás desde cero):
una hoja por folio en orden ascendente, sin la hoja residual `Sheet1`,
conservando estilos, celdas combinadas, alturas y el logo en TODAS las
hojas, incluso al desbordar materiales, actividades o comentarios; con
neutralización de fórmulas y eliminación de caracteres ilegales del XLSX.
"""

import hashlib
import io
from datetime import datetime, timezone
from math import ceil
from typing import Any, cast
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.actividad import ActividadMaterialOut, ActividadOut
from app.services import plantilla_reporte as pr
from app.services import plantilla_reporte_actividades as pa

ZONA = ZoneInfo("America/Mexico_City")
FECHA_BASE = datetime(2026, 10, 20, 9, 30, tzinfo=timezone.utc)


def _material(
    descripcion: str = "CABLE COAXIAL RG6",
    *,
    material_id: int = 1,
    u_m: str | None = "PZ",
    cantidad: float = 2.5,
    valor: float = 12.34,
) -> ActividadMaterialOut:
    return ActividadMaterialOut(
        material_id=material_id,
        codigo=None,
        descripcion=descripcion,
        u_m=u_m,
        cantidad_consumida=cantidad,
        valor_unitario=valor,
    )


def _actividad(
    *,
    folio_num: int = 1,
    fecha: datetime = FECHA_BASE,
    vehiculo: str = "Camioneta 12",
    clientes: str = "Cliente Uno",
    direcciones: str = "Av. Siempre Viva 123",
    actividades: list[str] | None = None,
    comentarios: str | None = None,
    materiales: list[ActividadMaterialOut] | None = None,
) -> ActividadOut:
    return ActividadOut(
        actividad_id=folio_num,
        folio=f"R{folio_num:04d}",
        folio_num=folio_num,
        seguimiento_id=1,
        equipo_id=1,
        equipo_nombre="EQUIPO UNO",
        fecha=fecha,
        usuario="qa",
        vehiculo=vehiculo,
        clientes=clientes,
        direcciones=direcciones,
        actividades=(
            actividades if actividades is not None else ["Actividad 1", "Actividad 2"]
        ),
        comentarios=comentarios,
        materiales=materiales if materiales is not None else [],
    )


def _generar(*actividades: ActividadOut) -> tuple[bytes, Workbook]:
    data = pa.generar_reporte_actividades(actividades=list(actividades), zona=ZONA)
    return data, cast(Workbook, load_workbook(io.BytesIO(data)))


def test_req_rep_007_distribucion_basica_de_campos_y_materiales() -> None:
    """REQ-REP-007: cabecera (FECHA, VEHÍCULO, NOMBRE, FOLIO, DIRECCION),
    tabla de materiales del snapshot (números y formato de VALOR UNITARIO),
    actividades una por fila envueltas, comentario y COLABORADORES."""
    materiales = [
        _material("CABLE COAXIAL RG6", material_id=1, cantidad=2.5, valor=12.34),
        _material(
            "CONECTOR F", material_id=2, u_m="BOLSA (500 PZ)", cantidad=7.0, valor=0.5
        ),
    ]
    realizadas = ["Tendido de cable principal", "Instalación de conectores"]
    _, wb = _generar(
        _actividad(
            materiales=materiales,
            actividades=realizadas,
            comentarios="Sin novedad",
        )
    )

    assert wb.sheetnames == ["R0001"]
    ws = wb["R0001"]
    assert ws["B8"].value == "20 / OCTUBRE / 2026"
    assert ws["E8"].value == "Camioneta 12"
    assert ws["B9"].value == "Cliente Uno"
    assert ws["I9"].value == "FOLIO: R0001"
    assert ws["B10"].value == "Av. Siempre Viva 123"

    assert ws["A13"].value == 2.5
    assert ws["B13"].value == "PZ"
    assert ws["C13"].value == "CABLE COAXIAL RG6"
    assert ws["I13"].value == 12.34
    assert ws["I13"].number_format == pa.FORMATO_VALOR_UNITARIO
    assert ws["A14"].value == 7.0
    assert ws["B14"].value == "BOLSA (500 PZ)"
    assert ws["C14"].value == "CONECTOR F"
    assert ws["I14"].value == 0.5
    assert ws["I14"].number_format == pa.FORMATO_VALOR_UNITARIO

    assert ws["A25"].value == "Tendido de cable principal"
    assert ws["A26"].value == "Instalación de conectores"
    for fila in (25, 26):
        celda = ws.cell(row=fila, column=1)
        assert celda.alignment.wrap_text is True
        assert celda.alignment.horizontal == "left"

    assert ws["C32"].value == "Sin novedad"
    assert ws["A36"].value == "COLABORADORES"


def test_req_rep_007_fecha_se_escribe_como_texto_sin_prefijo() -> None:
    """REQ-REP-007 / D-8: la FECHA es TEXTO `DD / MES / AAAA` en la zona de
    negocio, sin el prefijo «FECHA:» (la etiqueta ya está en A8)."""
    _, wb = _generar(_actividad())
    celda = wb["R0001"]["B8"]
    assert celda.data_type == "s"
    assert celda.value == "20 / OCTUBRE / 2026"
    assert not celda.value.startswith("FECHA:")


def test_req_rep_006_plantilla_en_disco_permanece_intacta() -> None:
    """REQ-REP-006: la generación carga la plantilla en memoria y jamás la
    modifica en disco (hash sha256 idéntico, incluso con desborde)."""
    ruta = pr.TEMPLATES_DIR / pa.PLANTILLA_ACTIVIDADES
    antes = hashlib.sha256(ruta.read_bytes()).hexdigest()
    _generar(_actividad(materiales=[_material(f"MATERIAL {i + 1}") for i in range(30)]))
    assert hashlib.sha256(ruta.read_bytes()).hexdigest() == antes


def test_req_rep_009_logo_en_todas_las_hojas_con_ancla_y_tamano_de_la_plantilla() -> None:
    """REQ-REP-009: el logo está en TODAS las hojas (una imagen por hoja),
    con el ancla y el tamaño de la imagen de la plantilla. Sin Pillow
    openpyxl descarta las imágenes: el test falla si falta la dependencia."""
    import PIL  # noqa: F401

    hoja_plantilla = cast(
        Worksheet,
        load_workbook(pr.TEMPLATES_DIR / pa.PLANTILLA_ACTIVIDADES)[pa.HOJA_PLANTILLA],
    )
    logo_plantilla = cast(Any, hoja_plantilla)._images[0]

    for n in (1, 3):
        _, wb = _generar(*[_actividad(folio_num=i) for i in range(1, n + 1)])
        assert len(wb.sheetnames) == n
        for nombre in wb.sheetnames:
            imagenes = cast(Any, wb[nombre])._images
            assert len(imagenes) == 1
            imagen = imagenes[0]
            ancla = imagen.anchor
            assert ancla._from.col == 6
            assert ancla._from.row == 0
            assert ancla._from.colOff == 560070
            assert ancla._from.rowOff == 22860
            assert ancla.ext.cx == 2266950
            assert ancla.ext.cy == 876300
            assert imagen.width == logo_plantilla.width
            assert imagen.height == logo_plantilla.height


def test_req_rep_009_hojas_ordenadas_por_folio_ascendente_sin_hoja_residual() -> None:
    """REQ-REP-009: una hoja por folio titulada con el folio, ordenada por
    folio ascendente aunque la entrada llegue en otro orden, y sin `Sheet1`."""
    _, wb = _generar(
        _actividad(folio_num=2, vehiculo="Vehiculo 2"),
        _actividad(folio_num=1),
    )
    assert wb.sheetnames == ["R0001", "R0002"]
    assert "Sheet1" not in wb.sheetnames
    assert wb["R0001"]["E8"].value == "Camioneta 12"
    assert wb["R0002"]["E8"].value == "Vehiculo 2"
    assert wb["R0001"]["I9"].value == "FOLIO: R0001"
    assert wb["R0002"]["I9"].value == "FOLIO: R0002"


def test_req_rep_008_actividad_sin_materiales_genera_tabla_vacia() -> None:
    """REQ-REP-008: una actividad con cero materiales genera la hoja con la
    tabla de materiales vacía y el resto del libro válido."""
    _, wb = _generar(_actividad(materiales=[]))
    ws = wb["R0001"]
    assert ws["A13"].value is None
    assert ws["B13"].value is None
    assert ws["C13"].value is None
    assert ws["I13"].value is None
    assert ws["B8"].value == "20 / OCTUBRE / 2026"
    assert ws["A25"].value == "Actividad 1"
    assert ws["A36"].value == "COLABORADORES"


def test_req_rep_008_desborde_de_materiales_inserta_filas_con_estilo_y_desplaza_colaboradores() -> None:
    """REQ-REP-008: más de 10 materiales insertan filas copiando estilo,
    altura y combinación C:H de la fila 13 y desplazan COLABORADORES."""
    once = [
        _material(
            f"MATERIAL {i + 1}",
            material_id=i + 1,
            cantidad=float(i + 1),
            valor=float(i + 1) / 2,
        )
        for i in range(11)
    ]
    _, wb = _generar(_actividad(materiales=once))
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert "C23:H23" in merges
    assert ws["A23"].value == 11
    assert ws["C23"].border.left.style == "thin"
    assert ws["C23"].border.bottom.style == "thin"
    assert ws.row_dimensions[23].height == ws.row_dimensions[13].height

    cuarenta = [
        _material(
            f"MATERIAL {i + 1}",
            material_id=i + 1,
            cantidad=float(i + 1),
            valor=float(i + 1) / 2,
        )
        for i in range(40)
    ]
    _, wb = _generar(_actividad(materiales=cuarenta))
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    # Filas insertadas con su combinación C:H; el último material (cantidad
    # 40) queda en la fila 13 + 39 = 52.
    assert "C42:H42" in merges
    assert "C52:H52" in merges
    assert ws["A52"].value == 40
    assert ws["C52"].value == "MATERIAL 40"
    # COLABORADORES se desplaza 30 filas (40 - 10 de desborde): 36 + 30 = 66.
    assert ws["A36"].value != "COLABORADORES"
    assert ws["A66"].value == "COLABORADORES"


def test_req_rep_008_desborde_de_actividades_inserta_filas_combinadas() -> None:
    """REQ-REP-008: más de 4 actividades insertan filas con la combinación
    A:I y el texto envuelto; 30 actividades de 500 caracteres caben íntegras
    con altura dentro de los límites de Excel."""
    cinco = [f"Actividad {i + 1}" for i in range(5)]
    _, wb = _generar(_actividad(actividades=cinco))
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert "A29:I29" in merges
    assert ws["A29"].value == "Actividad 5"
    assert ws["A29"].alignment.wrap_text is True

    treinta = [f"T{i + 1:02d} " + "A" * 496 for i in range(30)]
    assert len(treinta[29]) == 500
    _, wb = _generar(_actividad(actividades=treinta))
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert "A54:I54" in merges
    celda = ws["A54"]
    assert celda.value == treinta[29]
    assert celda.alignment.wrap_text is True
    assert 13.2 <= ws.row_dimensions[54].height <= 409.5


def test_req_rep_008_comentario_de_2000_caracteres_con_saltos_cumple_el_criterio_de_altura() -> None:
    """REQ-REP-008, criterio medible: suma de alturas de las 3 filas del área
    ≥ líneas estimadas × 13,2 pt con cpl conservador de 30 caracteres/línea
    para C:I, ninguna fila fuera de [13,2, 409,5] y el valor íntegro.

    Se usan 80 párrafos de 25 caracteres (2000 caracteres con saltos): con
    el cpl real del área (~86) cada párrafo ocupa una línea y el módulo
    reparte exactamente lo que exige el criterio conservador."""
    texto = "\n".join("P" * 25 for _ in range(80))
    assert len(texto.replace("\n", "")) == 2000
    _, wb = _generar(_actividad(comentarios=texto))
    ws = wb["R0001"]

    assert ws["C32"].value == texto
    assert ws["C32"].alignment.wrap_text is True

    lineas = sum(max(1, ceil(len(parrafo) / 30)) for parrafo in texto.split("\n"))
    assert lineas == 80
    alturas = [ws.row_dimensions[fila].height for fila in (32, 33, 34)]
    assert sum(alturas) >= lineas * 13.2
    for altura in alturas:
        assert 13.2 <= altura <= 409.5


def test_req_rep_008_comentario_de_12000_caracteres_amplia_el_area_y_desplaza_colaboradores() -> None:
    """REQ-REP-008: si la altura requerida supera la capacidad de las 3 filas
    del área se insertan filas, se extienden las combinaciones A:B y C:I,
    ninguna fila supera 409,5 pt y el texto queda íntegro; COLABORADORES se
    desplaza bajo el área ampliada."""
    texto = "P" * 12000
    _, wb = _generar(_actividad(comentarios=texto))
    ws = wb["R0001"]
    merges = {str(m) for m in ws.merged_cells.ranges}

    rango_valor = next(m for m in ws.merged_cells.ranges if str(m).startswith("C32:"))
    fila_fin = rango_valor.max_row
    insertadas = fila_fin - 34
    assert insertadas >= 1
    assert f"C32:I{fila_fin}" in merges
    assert f"A32:B{fila_fin}" in merges

    assert ws["C32"].value == texto
    for fila in range(32, fila_fin + 1):
        assert 13.2 <= ws.row_dimensions[fila].height <= 409.5

    assert ws["A36"].value is None
    fila_colaboradores = next(
        fila
        for fila in range(37, 60)
        if ws.cell(row=fila, column=1).value == "COLABORADORES"
    )
    assert fila_colaboradores == 36 + insertadas


def test_req_rep_008_descripcion_de_material_larga_se_envuelve_y_ajusta_altura() -> None:
    """REQ-REP-008: la descripción de material más larga que el ancho de su
    celda se envuelve y su fila crece (misma regla de alturas)."""
    descripcion = "D" * 200
    _, wb = _generar(_actividad(materiales=[_material(descripcion)]))
    ws = wb["R0001"]
    celda = ws["C13"]
    assert celda.value == descripcion
    assert celda.alignment.wrap_text is True
    assert ws.row_dimensions[13].height > 13.2


def test_req_rep_009_libro_con_100_hojas_ordenadas_y_logo_en_cada_una() -> None:
    """REQ-REP-009: 100 folios generan 100 hojas tituladas R0001..R0100 en
    orden ascendente, sin `Sheet1`, con una imagen de logo por hoja."""
    actividades = [
        _actividad(folio_num=n, actividades=[f"Actividad del folio {n}"])
        for n in range(1, 101)
    ]
    _, wb = _generar(*actividades)
    assert wb.sheetnames == [f"R{n:04d}" for n in range(1, 101)]
    assert "Sheet1" not in wb.sheetnames
    assert wb["R0001"]["I9"].value == "FOLIO: R0001"
    assert wb["R0100"]["I9"].value == "FOLIO: R0100"
    for nombre in wb.sheetnames:
        assert len(cast(Any, wb[nombre])._images) == 1


def test_req_rep_010_textos_tipo_formula_se_neutralizan_en_todos_los_campos() -> None:
    """REQ-REP-010 / C-10: todo texto que inicia (tras espacios) con `=`,
    `+`, `-` o `@` se prefija con un apóstrofe literal en cada campo libre y
    ninguna celda del libro queda como fórmula."""
    variantes = [
        ("=1+1", "'=1+1"),
        ("+10+5", "'+10+5"),
        ("- Instalación", "'- Instalación"),
        ("@SUM(A1:A9)", "'@SUM(A1:A9)"),
        ("  =1+1", "'  =1+1"),
        ("   +10+5", "'   +10+5"),
    ]
    for payload, esperado in variantes:
        actividad = _actividad(
            vehiculo=payload,
            clientes=payload,
            direcciones=payload,
            actividades=[payload, "Actividad normal"],
            comentarios=payload,
            materiales=[_material(payload)],
        )
        _, wb = _generar(actividad)
        for hoja in wb.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    assert celda.data_type != "f"
        ws = wb["R0001"]
        for coord in ("E8", "B9", "B10", "A25", "C32", "C13"):
            assert ws[coord].value == esperado
        assert ws["A26"].value == "Actividad normal"


def test_req_rep_010_texto_normal_no_recibe_apostrofe() -> None:
    """REQ-REP-010: un texto legítimo («Camioneta 12») se escribe tal cual,
    sin el apóstrofe de neutralización."""
    _, wb = _generar(_actividad(vehiculo="Camioneta 12"))
    assert wb["R0001"]["E8"].value == "Camioneta 12"


def test_req_rep_010_caracteres_ilegales_xlsx_se_eliminan_y_los_saltos_se_conservan() -> None:
    """REQ-REP-010: los caracteres ilegales del XLSX se eliminan antes de
    escribir (sin IllegalCharacterError) y los saltos de línea de los
    comentarios se conservan."""
    ilegales = "\x00\x01\x02\x07\x0b\x0c\x1f"
    actividad = _actividad(
        vehiculo=f"Camioneta{ilegales} 12",
        actividades=[f"tarea{ilegales}uno", "otra"],
        comentarios="primera\nsegunda",
    )
    _, wb = _generar(actividad)
    ws = wb["R0001"]
    assert ws["E8"].value == "Camioneta 12"
    assert ws["A25"].value == "tareauno"
    assert ws["C32"].value == "primera\nsegunda"
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str):
                assert not any(ch in celda.value for ch in ilegales)


def test_req_rep_007_unidad_de_medida_nula_deja_la_celda_vacia() -> None:
    """REQ-REP-007: una U.M. nula del snapshot no escribe nada en la celda
    (queda vacía, jamás la cadena "None")."""
    materiales = [
        _material("CABLE", material_id=1, u_m="PZ", cantidad=1.0, valor=2.0),
        _material("CONECTOR", material_id=2, u_m=None, cantidad=3.0, valor=4.0),
    ]
    _, wb = _generar(_actividad(materiales=materiales))
    ws = wb["R0001"]
    assert ws["B13"].value == "PZ"
    assert ws["B14"].value is None


def test_req_rep_010_nombre_de_archivo_con_una_hoja_usa_el_folio() -> None:
    """REQ-REP-010 / C-5: con una sola actividad el adjunto se nombra con el
    folio."""
    assert (
        pa.nombre_archivo_actividades(actividades=[_actividad()], zona=ZONA)
        == "reporte-actividades-R0001.xlsx"
    )


def test_req_rep_010_nombre_de_archivo_consolidado_usa_la_fecha_maxima_en_la_zona() -> None:
    """REQ-REP-010 / C-5: con dos o más hojas el adjunto usa la fecha de la
    actividad más reciente en la zona de negocio (2026-10-21T02:00Z es
    2026-10-20 en America/Mexico_City, UTC-6)."""
    temprana = _actividad(fecha=datetime(2026, 10, 20, 9, 0, tzinfo=timezone.utc))
    tardia = _actividad(
        folio_num=2, fecha=datetime(2026, 10, 21, 2, 0, tzinfo=timezone.utc)
    )
    assert (
        pa.nombre_archivo_actividades(actividades=[temprana, tardia], zona=ZONA)
        == "reporte-actividades-2026-10-20.xlsx"
    )


def test_req_rep_010_lista_vacia_lanza_value_error() -> None:
    """REQ-REP-010: sin actividades no hay libro ni nombre que construir."""
    with pytest.raises(ValueError):
        pa.generar_reporte_actividades(actividades=[], zona=ZONA)
    with pytest.raises(ValueError):
        pa.nombre_archivo_actividades(actividades=[], zona=ZONA)
