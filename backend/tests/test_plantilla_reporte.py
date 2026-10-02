"""Servicio PlantillaReporte — REQ-P2-REP-001/002/004/006/007.

Verifica que el XLSX se obtiene INYECTANDO datos sobre las plantillas
corporativas (jamás desde cero), conservando estilos, celdas combinadas y el
logo embebido, incluso al desbordar las filas de la plantilla.
"""

import hashlib
import io
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

from app.services import plantilla_reporte as pr
from app.services.plantilla_reporte import LineaActa

# formato -> (capacidad base, primera fila de datos, fila de inserción)
FORMATOS = {
    "normal": (16, 10, 26),
    "extenso": (20, 9, 29),
}


def _plantilla(formato: str):
    return load_workbook(pr.TEMPLATES_DIR / pr._PLANTILLAS[formato]).active


def _generar(formato: str, n: int, tipo: str = "ENTREGA") -> tuple[bytes, object]:
    lineas = [LineaActa(i + 1, f"MATERIAL {i}", "PZ") for i in range(n)]
    data = pr.generar_acta(
        tipo=tipo, formato=formato, fecha=date(2026, 10, 2), lineas=lineas
    )
    return data, load_workbook(io.BytesIO(data)).active


def test_pillow_instalado_requisito_para_conservar_el_logo() -> None:
    """Sin Pillow openpyxl descarta silenciosamente las imágenes: el test
    FALLA si la dependencia no está instalada (REQ de T2.1)."""
    import PIL  # noqa: F401

    _, ws = _generar("normal", 1)
    assert len(ws._images) == 1


@pytest.mark.parametrize("formato", ["normal", "extenso"])
def test_las_plantillas_corporativas_no_se_modifican(formato: str) -> None:
    ruta = pr.TEMPLATES_DIR / pr._PLANTILLAS[formato]
    antes = hashlib.sha256(ruta.read_bytes()).hexdigest()
    _generar(formato, 30)
    assert hashlib.sha256(ruta.read_bytes()).hexdigest() == antes


@pytest.mark.parametrize("formato", ["normal", "extenso"])
def test_inyeccion_sobre_plantilla_conserva_estilo_merges_y_logo(
    formato: str,
) -> None:
    capacidad, primera, _ = FORMATOS[formato]
    data, ws = _generar(formato, 3)
    orig = _plantilla(formato)

    # Viene de la plantilla (no de cero): conserva su encabezado y firmas.
    textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("NOMBRE" in t and "FIRMA" in t for t in textos)
    assert ws["A2"].value == orig["A2"].value  # ENTREGA no cambia el título

    # Datos inyectados en CANT. (A), DESCRIPCION (B:D) y U.M. (E).
    assert ws.cell(row=primera, column=1).value == 1
    assert ws.cell(row=primera, column=2).value == "MATERIAL 0"
    assert ws.cell(row=primera, column=5).value == "PZ"
    assert ws["A5"].value == "FECHA: 02 / OCTUBRE / 2026"

    # Merges idénticos a la plantilla; estilo del encabezado intacto.
    assert {str(r) for r in ws.merged_cells.ranges} == {
        str(r) for r in orig.merged_cells.ranges
    }
    celda_hdr = "A9" if formato == "normal" else "A8"
    for coord in ("A2", "A5", celda_hdr):
        a, b = orig[coord], ws[coord]
        assert (a.font.sz, a.font.b, a.fill.fgColor.rgb, a.border.left.style) == (
            b.font.sz, b.font.b, b.fill.fgColor.rgb, b.border.left.style,
        )  # fmt: skip

    # Logo embebido conservado (archivo de media + ancla).
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert any(n.startswith("xl/media/") and n.endswith(".png") for n in z.namelist())
    assert len(ws._images) == 1
    ancla = ws._images[0].anchor
    oancla = orig._images[0].anchor
    assert (ancla._from.col, ancla._from.row) == (oancla._from.col, oancla._from.row)
    assert (ancla.ext.width, ancla.ext.height) == (oancla.ext.width, oancla.ext.height)


@pytest.mark.parametrize("formato", ["normal", "extenso"])
@pytest.mark.parametrize("exceso", [1, 5, 25])
def test_desborde_inserta_filas_y_desplaza_firmas_merges_y_logo(
    formato: str, exceso: int
) -> None:
    capacidad, primera, insercion = FORMATOS[formato]
    n = capacidad + exceso
    data, ws = _generar(formato, n)
    orig = _plantilla(formato)

    # Todas las líneas quedan escritas, sin huecos.
    for i in range(n):
        assert ws.cell(row=primera + i, column=2).value == f"MATERIAL {i}"
        assert ws.cell(row=primera + i, column=1).value == i + 1

    # Merges: los de la plantilla bajo la inserción se desplazan `exceso`
    # filas y cada fila nueva tiene su combinación B:D.
    esperados = set()
    for r in orig.merged_cells.ranges:
        cr = CellRange(str(r))
        if cr.min_row >= insercion:
            cr.shift(row_shift=exceso)
        esperados.add(cr.coord)
    for k in range(exceso):
        esperados.add(f"B{insercion + k}:D{insercion + k}")
    assert {str(r) for r in ws.merged_cells.ranges} == esperados

    # Los bloques de firma/colaboradores se desplazaron intactos.
    for fila in range(insercion, orig.max_row + 1):
        for col in range(1, orig.max_column + 1):
            a = orig.cell(row=fila, column=col)
            b = ws.cell(row=fila + exceso, column=col)
            if isinstance(a.value, str):
                assert b.value == a.value
        assert ws.row_dimensions[fila + exceso].height == orig.row_dimensions[fila].height

    # Estilo de las filas nuevas = estilo de las filas de datos (bordes).
    for fila in range(primera, primera + n):
        for col in (1, 2, 5):
            c = ws.cell(row=fila, column=col)
            assert c.border.bottom.style == "thin"
            assert c.alignment.horizontal == "center"

    # El logo sigue embebido y anclado en su posición original.
    assert len(ws._images) == 1
    assert ws._images[0].anchor._from.row == orig._images[0].anchor._from.row
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert any(n.startswith("xl/media/") for n in z.namelist())


@pytest.mark.parametrize("formato", ["normal", "extenso"])
def test_acta_de_consumo_ajusta_el_titulo(formato: str) -> None:
    _, ws = _generar(formato, 2, tipo="CONSUMO")
    assert ws["A2"].value == "ACTA DE CONSUMO"
    # El título conserva el estilo de la plantilla.
    assert ws["A2"].font.sz == _plantilla(formato)["A2"].font.sz
    assert ws["A2"].font.b is True


def test_formato_fecha_en_espanol() -> None:
    assert pr.formatear_fecha(date(2026, 9, 28)) == "FECHA: 28 / SEPTIEMBRE / 2026"
    assert pr.formatear_fecha(date(2026, 1, 5)) == "FECHA: 05 / ENERO / 2026"


@pytest.mark.parametrize(
    "valor,esperado",
    [
        ("=cmd|' /C calc'!A0", "'=cmd|' /C calc'!A0"),
        ("+1+1", "'+1+1"),
        ("-2", "'-2"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("  =HYPERLINK()", "'  =HYPERLINK()"),
        ("Cable normal", "Cable normal"),
        (7, 7),
        (None, None),
    ],
)
def test_sanitizacion_anti_formula(valor, esperado) -> None:
    assert pr.sanitizar_celda(valor) == esperado


def test_celda_maliciosa_se_escribe_como_texto_no_como_formula() -> None:
    lineas = [LineaActa(1, "=cmd|' /C calc'!A0", "@x")]
    data = pr.generar_acta(
        tipo="ENTREGA", formato="normal", fecha=date(2026, 10, 2), lineas=lineas
    )
    ws = load_workbook(io.BytesIO(data)).active
    for coord in ("B10", "E10"):
        assert ws[coord].data_type == "s"
        assert ws[coord].value.startswith("'")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert b"<f>" not in z.read("xl/worksheets/sheet1.xml")


def test_nombre_de_archivo_sanitizado_y_solo_constantes() -> None:
    assert pr.nombre_archivo_acta("ENTREGA", "normal", date(2026, 10, 2)) == (
        "acta-entrega-normal-2026-10-02.xlsx"
    )
    peligroso = pr.sanitizar_nombre_archivo('../../etc/passwd"\r\nX: y')
    assert peligroso == "etc_passwd___X__y"
    assert not any(c in peligroso for c in '/\\"\r\n:')
    assert pr.sanitizar_nombre_archivo("...") == "acta"


def test_formato_desconocido_no_toca_el_sistema_de_archivos() -> None:
    """Las rutas son constantes del servidor: un formato arbitrario (p. ej.
    un intento de path traversal) se rechaza sin construir ninguna ruta."""
    with pytest.raises(ValueError):
        pr.generar_acta(
            tipo="ENTREGA",
            formato="../../etc/passwd",  # type: ignore[arg-type]
            fecha=date(2026, 10, 2),
            lineas=[],
        )
